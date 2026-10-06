import os
import subprocess

import fitz
from docx import Document

print(" CONTRACT FILE SERVICE LOADED ")

def extract_text_from_pdf(file_path):
    doc = fitz.open(file_path)

    text = ""

    for page in doc:
        page_text = page.get_text()

        if page_text:
            text += page_text + "\n"

    doc.close()

    return text


def extract_text_from_docx(file_path):
    document = Document(file_path)

    text_parts = []

    # Normal paragraphs
    for paragraph in document.paragraphs:
        if paragraph.text.strip():
            text_parts.append(paragraph.text)

    # Tables
    for table in document.tables:
        for row in table.rows:
            row_text = []

            for cell in row.cells:
                cell_text = cell.text.strip()

                if cell_text:
                    row_text.append(cell_text)

            if row_text:
                text_parts.append(" | ".join(row_text))

    return "\n".join(text_parts)


def extract_text_from_doc(file_path):
    try:
        result = subprocess.run(
            ["antiword", file_path],
            stdout=subprocess.PIPE,
            stderr=subprocess.PIPE,
            check=True,
        )

        try:
            return result.stdout.decode("utf-8")

        except UnicodeDecodeError:
            return result.stdout.decode(
                "latin-1",
                errors="replace"
            )

    except FileNotFoundError:
        raise RuntimeError(
            "DOC file processing requires Antiword on the server."
        )

    except subprocess.CalledProcessError as error:
        error_message = error.stderr.decode(
            "utf-8",
            errors="replace"
        )

        raise RuntimeError(
            f"Failed to extract text from DOC file: {error_message}"
        )


def extract_text_from_file(file_path):
    extension = os.path.splitext(file_path)[1].lower()

    if extension == ".pdf":
        return extract_text_from_pdf(file_path)

    if extension == ".docx":
        return extract_text_from_docx(file_path)

    if extension == ".doc":
        return extract_text_from_doc(file_path)

    raise ValueError(
        "Unsupported file format. Please upload PDF, DOC, or DOCX."
    )